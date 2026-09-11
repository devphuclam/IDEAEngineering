from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import build_report as base


ROOT = Path(r"C:\Users\TD-999\Research\Projects\IDEA\IDEAEngineering")
HUMAN = Path(r"C:\Users\TD-999\Research\Projects\IDEA\Human\IDEAEngineering")
REFERENCE = HUMAN / "IDEA_DDM_review.docx"
OUTPUT = HUMAN / "IDEA-DDM-bao-cao-kien-truc-nghiep-vu-trinh-review-v2.docx"
LOGO = ROOT / "logo-idea.png"
WORK = ROOT / ".tmp-artifact" / "architecture-report-v2"
FIG_DIR = WORK / "figures"

INSTANCE = ROOT / "docs" / "product" / "instances" / "idea-engineering"
EVIDENCE = INSTANCE / "evidence" / "IE-VEV-ARCH-VIEW-002"
GALLERY = EVIDENCE / "index.html"
DOC03 = INSTANCE / "DOC-03-business-requirements.md"
DOC04 = INSTANCE / "DOC-04-software-requirements-specification.md"
DOC05 = INSTANCE / "DOC-05-architecture-description.md"
DOC06 = INSTANCE / "DOC-06-data-integration-and-migration-specification.md"


FIGURE_ALT = {
    "figure-01-users.png": "Sơ đồ các nhóm người dùng IDEA DDM và công việc chính từ soạn tài liệu đến tra cứu hồ sơ đã phát hành.",
    "figure-02-blocks.png": "Sơ đồ sáu khối chính gồm Web, Desktop, Server, PostgreSQL, kho file và bộ xử lý định dạng, kèm ranh giới trách nhiệm.",
    "figure-03-checkout-reference-choice.png": "Sơ đồ quyết định dùng Checkout khi cần sửa hoặc Reference khi chỉ cần tham khảo.",
    "figure-04-checkout-sequence.png": "Sơ đồ trình tự Checkout từ chọn tài liệu, kiểm tra quyền và trạng thái đến tạo bản ghi giữ sửa và tải file vào Workspace.",
    "figure-05-reference-handling.png": "Sơ đồ các cách xử lý an toàn khi file Reference được giữ nguyên, bị sửa hoặc đã dựa trên bản cũ.",
    "figure-06-checkin-outcomes.png": "Sơ đồ Check-in với bốn kết quả: có thay đổi, không thay đổi, bị từ chối trước khi ghi hoặc kết quả chưa xác định do mất kết nối.",
    "figure-07-stale-recovery.png": "Sơ đồ xử lý bản làm việc đã cũ và thao tác Check-in bị gián đoạn mà vẫn giữ an toàn file cục bộ.",
    "figure-08-document-lifecycle.png": "Sơ đồ trạng thái tài liệu từ Start, In Work, Under Review đến Released, gồm nhánh Reject và lưu ý Checkout không đổi trạng thái workflow.",
    "figure-09-staged-release.png": "Sơ đồ phát hành theo phạm vi P-100, trong đó cụm bơm có thể Released trong khi tủ điện và hồ sơ toàn máy vẫn In Work.",
    "figure-10-rbac-model.png": "Sơ đồ RBAC cho thấy Principal, Role Definition và Scope tạo thành Role Assignment; quyền hiệu lực còn phải qua điều kiện nghiệp vụ.",
    "figure-11-project-permission.png": "Sơ đồ ví dụ Linh được cấp tài khoản, tham gia Project P-100, vào Group Cơ khí và nhận quyền qua Role Assignment tại đúng phạm vi Project.",
    "figure-12-admin-separation.png": "Sơ đồ tách trách nhiệm giữa quản trị tài khoản, quản trị Project, quản trị Role, quản trị cấu hình sản phẩm, tra cứu Audit và Super Administrator.",
    "figure-13-auth-decision.png": "Sơ đồ trình tự kiểm tra một yêu cầu từ phiên đăng nhập, thành viên Project và Group, Role Assignment, Permission đến business gate và Audit.",
    "figure-14-operation.png": "Sơ đồ triển khai logic, backup độc lập, ranh giới bảo mật và đường mở rộng của IDEA DDM mà không chọn stack hay phiên bản công nghệ.",
}


def _save(img: Image.Image, name: str) -> Path:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    path = FIG_DIR / name
    img.save(path, format="PNG", dpi=(180, 180), optimize=True)
    return path


def _label(draw: ImageDraw.ImageDraw, xy, text: str, size=25, color="#526171", bold=False, anchor="la"):
    draw.text(xy, text, font=base.font(size, bold), fill=color, anchor=anchor)


def _panel(draw: ImageDraw.ImageDraw, xy, title: str, fill="#F6F8FA", outline="#D6E0EA"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=fill, outline=outline, width=3)
    draw.text((x1 + 22, y1 + 18), title, font=base.font(30, True), fill="#0F2747")


def _sequence_lanes(draw: ImageDraw.ImageDraw, names: list[str], top=205, bottom=1010):
    left, right = 110, 1890
    step = (right - left) / (len(names) - 1)
    xs = []
    for idx, name in enumerate(names):
        x = int(left + idx * step)
        xs.append(x)
        draw.rounded_rectangle((x - 130, top - 45, x + 130, top + 45), radius=14, fill="#EAF3FA", outline="#1F6FB2", width=3)
        draw.text((x, top), name, font=base.font(25, True), fill="#0F2747", anchor="mm")
        draw.line((x, top + 46, x, bottom), fill="#B8C5D1", width=3)
    return xs


def _message(draw: ImageDraw.ImageDraw, x1, x2, y, text, color="#526171", dashed=False):
    if dashed:
        span = abs(x2 - x1)
        count = max(1, int(span / 30))
        direction = 1 if x2 > x1 else -1
        for i in range(count):
            if i % 2 == 0:
                a = x1 + direction * (span * i / count)
                b = x1 + direction * (span * min(i + 1, count) / count)
                draw.line((a, y, b, y), fill=color, width=4)
        base.arrow(draw, (x2 - (15 if x2 > x1 else -15), y), (x2, y), color=color, width=4, head=15)
    else:
        base.arrow(draw, (x1, y), (x2, y), color=color, width=4, head=15)
    mid = (x1 + x2) / 2
    lines = base.wrap_text(draw, text, base.font(22), max(230, abs(x2 - x1) - 30))
    for i, line in enumerate(lines[:2]):
        draw.text((mid, y - 30 - (len(lines[:2]) - 1 - i) * 24), line, font=base.font(22), fill=color, anchor="mm")


def make_figures() -> dict[str, Path]:
    figs: dict[str, Path] = {}

    # 1. Users and work
    img, d = base.diagram_base("Ai sử dụng IDEA DDM?", "Mỗi nhóm nhìn cùng một tài liệu nhưng thực hiện trách nhiệm khác nhau")
    users = [
        ("Kỹ sư / người soạn", "Tìm tài liệu, Checkout, sửa và Check-in"),
        ("Người review", "Xem đúng bản được gửi và ghi quyết định"),
        ("Người Release", "Xác nhận đúng phạm vi hồ sơ trước khi phát hành"),
        ("Quản trị viên", "Cấp tài khoản, quyền, Project và cấu hình theo thẩm quyền"),
        ("Người tra cứu", "Mở đúng bản được phép sử dụng và xem lịch sử"),
    ]
    for i, (title, subtitle) in enumerate(users):
        y = 220 + i * 160
        base.box(d, (90, y, 560, y + 115), title, subtitle, fill="#F6F9FC", title_size=30, body_size=23)
        base.arrow(d, (560, y + 58), (830, 560), color="#7B8FA3", width=4, head=15)
    base.box(d, (830, 400, 1260, 720), "IDEA DDM", "Quản lý danh tính tài liệu, nội dung, cấu trúc, quyền, lịch sử và bản phát hành", fill="#EAF3FA", title_size=42, body_size=30)
    outcomes = [
        ("Đúng tài liệu", "Biết bản nào đang sửa và bản nào được dùng"),
        ("Đúng người", "Chỉ người đủ quyền mới thực hiện thao tác"),
        ("Đúng bộ hồ sơ", "Release đúng phạm vi đã xem xét"),
        ("Lấy lại được", "Tái tạo đúng bộ đã phát hành trước đây"),
    ]
    for i, (title, subtitle) in enumerate(outcomes):
        y = 245 + i * 185
        base.box(d, (1480, y, 1910, y + 125), title, subtitle, fill="#E8F5EE", outline="#2E7D5B", title_size=30, body_size=23)
        base.arrow(d, (1260, 560), (1480, y + 63), color="#2E7D5B", width=4, head=15)
    figs["users"] = _save(img, "figure-01-users.png")

    # 2. Six blocks
    img, d = base.diagram_base("Sáu khối chính của IDEA DDM", "Giao diện nhận thao tác; Server quyết định; dữ liệu và file được lưu ở đúng nơi")
    base.box(d, (90, 250, 480, 440), "Web", "Tìm kiếm, xem, review và màn hình quản trị được cấp quyền", fill="#F6F9FC", title_size=34, body_size=25)
    base.box(d, (90, 650, 480, 840), "Desktop", "Quản lý Workspace và mở file bằng Office/CAD trên Windows", fill="#F6F9FC", title_size=34, body_size=25)
    base.box(d, (720, 350, 1240, 740), "IDEA Server", "Xác thực, RBAC, Checkout, Check-in, workflow, Release và Audit", fill="#EAF3FA", title_size=42, body_size=30)
    base.box(d, (1510, 220, 1910, 390), "PostgreSQL", "Metadata, trạng thái, quyền và giao dịch", fill="#E8F5EE", outline="#2E7D5B", title_size=33, body_size=24)
    base.box(d, (1510, 475, 1910, 645), "Kho file", "Nội dung bất biến và dấu kiểm tra toàn vẹn", fill="#E8F5EE", outline="#2E7D5B", title_size=33, body_size=24)
    base.box(d, (1510, 730, 1910, 900), "Bộ xử lý định dạng", "Tạo preview hoặc bản dẫn xuất trong vùng tách biệt", fill="#FFF4DB", outline="#A96C12", title_size=31, body_size=23)
    base.arrow(d, (480, 345), (720, 470), label="HTTPS")
    base.arrow(d, (480, 745), (720, 625), label="HTTPS")
    base.arrow(d, (1240, 445), (1510, 305), label="giao dịch")
    base.arrow(d, (1240, 545), (1510, 560), label="nội dung")
    base.arrow(d, (1240, 650), (1510, 815), label="job giới hạn")
    _label(d, (1000, 1015), "Báo cáo mô tả trách nhiệm của các khối, không đề xuất stack, phiên bản hoặc nhà cung cấp khác.", size=24, color="#0F2747", bold=True, anchor="mm")
    figs["blocks"] = _save(img, "figure-02-blocks.png")

    # 3. Checkout or Reference decision
    img, d = base.diagram_base("Chọn Checkout hay Reference", "Người dùng chọn theo mục đích; hệ thống không tự suy rộng sang cả cây tài liệu")
    base.box(d, (720, 205, 1280, 355), "Bạn cần làm gì với tài liệu?", "Chọn đúng phạm vi trước khi tải file", fill="#EAF3FA", title_size=36, body_size=25)
    base.arrow(d, (850, 355), (470, 520), label="Cần sửa")
    base.arrow(d, (1150, 355), (1530, 520), label="Chỉ tham khảo")
    base.box(d, (160, 520, 780, 760), "Checkout", "Giữ quyền sửa cho đúng người và đúng Workspace. Người khác không thể Checkout cùng tài liệu khi quyền giữ còn hiệu lực.", fill="#EAF3FA", title_size=40, body_size=27)
    base.box(d, (1220, 520, 1840, 760), "Reference", "Tải đúng bản để xem hoặc đối chiếu. Reference không có quyền Check-in vào tài liệu gốc.", fill="#F6F8FA", outline="#7B8FA3", title_size=40, body_size=27)
    base.box(d, (250, 855, 690, 1015), "Phạm vi rõ ràng", "Chỉ các tài liệu được xác nhận mới được giữ sửa", fill="#E8F5EE", outline="#2E7D5B", title_size=30, body_size=23)
    base.box(d, (1310, 855, 1750, 1015), "Không khóa tài liệu", "Nhiều người có thể lấy Reference cùng lúc", fill="#E8F5EE", outline="#2E7D5B", title_size=30, body_size=23)
    base.arrow(d, (470, 760), (470, 855), color="#2E7D5B")
    base.arrow(d, (1530, 760), (1530, 855), color="#2E7D5B")
    figs["choice"] = _save(img, "figure-03-checkout-reference-choice.png")

    # 4. Checkout sequence
    img, d = base.diagram_base("Checkout diễn ra như thế nào", "Không mở file trước khi Server xác nhận quyền, trạng thái và đúng bản nguồn")
    xs = _sequence_lanes(d, ["Kỹ sư", "Web hoặc Desktop", "IDEA Server", "Workspace"])
    y = 330
    _message(d, xs[0], xs[1], y, "Chọn tài liệu và phạm vi")
    y += 115
    _message(d, xs[1], xs[2], y, "Yêu cầu Checkout đúng Generation")
    y += 115
    _message(d, xs[2], xs[2] + 1, y, "Kiểm tra tài khoản, RBAC, trạng thái và quyền giữ", color="#A96C12")
    y += 115
    _message(d, xs[2], xs[1], y, "Xác nhận Reservation và danh sách file", color="#2E7D5B", dashed=True)
    y += 115
    _message(d, xs[1], xs[3], y, "Tải file, kiểm tra digest và ghi manifest")
    y += 115
    _message(d, xs[3], xs[0], y, "Mở file làm việc", color="#2E7D5B", dashed=True)
    _label(d, (1000, 1030), "Nếu một tài liệu trong phạm vi không hợp lệ, phần yêu cầu giữ sửa bị từ chối và không mở nhầm file.", size=24, color="#A43B3B", bold=True, anchor="mm")
    figs["checkout"] = _save(img, "figure-04-checkout-sequence.png")

    # 5. Reference handling
    img, d = base.diagram_base("Reference được xử lý như thế nào", "File tham khảo không thể ghi ngược vào tài liệu gốc nếu chưa có Checkout hợp lệ")
    base.box(d, (710, 205, 1290, 355), "Reference trong Workspace", "Gắn đúng tài liệu, Generation và digest", fill="#F6F8FA", outline="#7B8FA3", title_size=37, body_size=25)
    branches = [
        ("Không sửa", "Tiếp tục xem hoặc xóa bản cục bộ", "#E8F5EE", "#2E7D5B", 120),
        ("Đã sửa, bản nguồn vẫn hiện hành", "Có thể xin Checkout đúng Generation rồi giữ lại file cục bộ", "#EAF3FA", "#1F6FB2", 775),
        ("Đã sửa, bản nguồn đã mới hơn", "Không Check-in trực tiếp. Giữ file để so sánh, áp dụng lại có chủ đích, tạo bản sao hoặc bỏ sau xác nhận", "#FCECEC", "#A43B3B", 1430),
    ]
    for title, subtitle, fill, outline, x in branches:
        base.arrow(d, (1000, 355), (x + 225, 520), color=outline, width=4, head=15)
        base.box(d, (x, 520, x + 450, 825), title, subtitle, fill=fill, outline=outline, title_size=31, body_size=25)
    _label(d, (1000, 930), "Mọi nhánh đều giữ tài liệu gốc nguyên vẹn cho đến khi có một Checkout và Check-in hợp lệ.", size=26, color="#0F2747", bold=True, anchor="mm")
    figs["reference"] = _save(img, "figure-05-reference-handling.png")

    # 6. Check-in outcomes
    img, d = base.diagram_base("Một lần Check-in có thể kết thúc ra sao", "Chỉ giao dịch thành công mới thay đổi dữ liệu chính thức; lỗi không để lại kết quả dở dang")
    base.box(d, (650, 205, 1350, 365), "Phạm vi Check-in đã xác nhận", "Server kiểm tra lại quyền, Checkout, expected Generation và toàn bộ file", fill="#EAF3FA", title_size=36, body_size=25)
    outcomes = [
        ("Có thay đổi", "Tạo Generation mới, tăng Version và kết thúc Checkout", "#E8F5EE", "#2E7D5B"),
        ("Không thay đổi", "Không tăng Version hoặc Generation nhưng vẫn kết thúc Checkout", "#E8F5EE", "#2E7D5B"),
        ("Bị từ chối trước commit", "Không công bố bản mới, không ghi một phần và giữ file cục bộ", "#FCECEC", "#A43B3B"),
        ("Mất kết nối quanh commit", "Tra cứu lại cùng OperationId trước khi gửi lại hoặc kết luận", "#FFF4DB", "#A96C12"),
    ]
    for i, (title, subtitle, fill, outline) in enumerate(outcomes):
        x = 70 + i * 485
        base.arrow(d, (1000, 365), (x + 215, 535), color=outline, width=4, head=15)
        base.box(d, (x, 535, x + 430, 850), title, subtitle, fill=fill, outline=outline, title_size=30, body_size=24)
    _label(d, (1000, 970), "Check-in nhiều tài liệu dùng một kết quả all-or-none: hoặc toàn bộ phạm vi hợp lệ được ghi, hoặc không tài liệu nào được công bố.", size=24, color="#0F2747", bold=True, anchor="mm")
    figs["checkin"] = _save(img, "figure-06-checkin-outcomes.png")

    # 7. Stale and interrupted recovery
    img, d = base.diagram_base("Xử lý bản cũ và Check-in bị gián đoạn", "Hai tình huống khác nhau nhưng đều không được tự ghi đè hoặc tự đoán kết quả")
    _panel(d, (70, 210, 965, 995), "A. Bản làm việc đã cũ")
    base.box(d, (165, 310, 870, 460), "Expected Generation khác Working Head", "Server phát hiện trước hoặc ngay lúc commit", fill="#FCECEC", outline="#A43B3B", title_size=31, body_size=24)
    base.arrow(d, (520, 460), (520, 585), color="#A43B3B")
    base.box(d, (165, 585, 870, 770), "Không ghi đè và không tự merge", "Giữ file cục bộ. Người dùng lấy bản mới, so sánh rồi áp dụng lại thay đổi có chủ đích.", fill="#FFF4DB", outline="#A96C12", title_size=31, body_size=24)
    base.arrow(d, (520, 770), (520, 875), color="#2E7D5B")
    _label(d, (520, 920), "Tài liệu chính thức không đổi", size=28, color="#2E7D5B", bold=True, anchor="mm")
    _panel(d, (1035, 210, 1930, 995), "B. Không nhận được phản hồi Check-in")
    base.box(d, (1130, 310, 1835, 460), "Không gửi một thao tác mới", "Client hỏi trạng thái bằng đúng OperationId", fill="#FFF4DB", outline="#A96C12", title_size=31, body_size=24)
    branches = [
        ("Đã commit", "Nhận lại đúng kết quả cũ", "#E8F5EE", "#2E7D5B"),
        ("Chưa commit", "Tiếp tục phần tải còn thiếu", "#EAF3FA", "#1F6FB2"),
        ("Chưa xác định", "Giữ thao tác để đối soát", "#FCECEC", "#A43B3B"),
    ]
    for i, (title, subtitle, fill, outline) in enumerate(branches):
        x = 1070 + i * 285
        base.arrow(d, (1480, 460), (x + 130, 605), color=outline, width=4, head=15)
        base.box(d, (x, 605, x + 260, 840), title, subtitle, fill=fill, outline=outline, title_size=27, body_size=22)
    figs["recovery"] = _save(img, "figure-07-stale-recovery.png")

    # 8. Lifecycle
    img, d = base.diagram_base("Vòng đời tài liệu", "Checkout kiểm soát quyền sửa; workflow kiểm soát trạng thái xét duyệt và phát hành")
    stages = [
        ("Start", "Chưa có bản làm việc", "#F6F8FA", "#7B8FA3"),
        ("In Work", "Đang soạn và Check-in", "#EAF3FA", "#1F6FB2"),
        ("Under Review", "Đúng Generation đang được xét", "#FFF4DB", "#A96C12"),
        ("Released", "Đã phát hành đúng phạm vi", "#E8F5EE", "#2E7D5B"),
    ]
    boxes = []
    for i, (title, subtitle, fill, outline) in enumerate(stages):
        x = 80 + i * 485
        boxes.append((x, x + 400))
        base.box(d, (x, 375, x + 400, 650), title, subtitle, fill=fill, outline=outline, title_size=38, body_size=25)
        if i:
            label = ["Check-in đầu tiên", "Submit", "Release đủ điều kiện"][i - 1]
            base.arrow(d, (boxes[i - 1][1], 510), (x, 510), color=outline, label=label, label_y_offset=-55)
    # Reject returns the exact Generation under review to In Work. It must not
    # look like a transition out of Released.
    base.arrow(d, (1250, 650), (1000, 790), color="#A43B3B", label="Reject", label_y_offset=15)
    base.arrow(d, (1000, 790), (765, 650), color="#A43B3B")
    _label(d, (1000, 925), "Checkout, gia hạn hoặc kết thúc Checkout không tự chuyển trạng thái workflow.", size=28, color="#0F2747", bold=True, anchor="mm")
    figs["lifecycle"] = _save(img, "figure-08-document-lifecycle.png")

    # 9. Staged release P-100
    img, d = base.diagram_base("Phát hành theo đúng phạm vi P-100", "Một cụm hoàn tất có thể được phát hành mà không kéo các phần chưa hoàn tất sang Released")
    base.box(d, (730, 210, 1270, 355), "Hồ sơ P-100", "Ba phạm vi có tiến độ độc lập", fill="#EAF3FA", title_size=38, body_size=25)
    items = [
        ("Cụm bơm", "Đã review, đủ phụ thuộc bắt buộc", "Released", "#E8F5EE", "#2E7D5B", 110),
        ("Tủ điện", "Thiết kế vẫn tiếp tục", "In Work", "#EAF3FA", "#1F6FB2", 780),
        ("Hồ sơ toàn máy", "Chưa đủ tài liệu để phát hành", "In Work", "#EAF3FA", "#1F6FB2", 1450),
    ]
    for title, subtitle, state, fill, outline, x in items:
        base.arrow(d, (1000, 355), (x + 220, 520), color=outline, width=4, head=15)
        base.box(d, (x, 520, x + 440, 795), title, f"{subtitle}. Trạng thái: {state}.", fill=fill, outline=outline, title_size=35, body_size=25)
    base.box(d, (260, 900, 1740, 1030), "Release Record của cụm bơm", "Chỉ pin đúng Generation, Structure Snapshot, phụ thuộc bắt buộc và bằng chứng thuộc phạm vi cụm bơm.", fill="#E8F5EE", outline="#2E7D5B", title_size=30, body_size=24)
    figs["staged_release"] = _save(img, "figure-09-staged-release.png")

    # 10. RBAC model
    img, d = base.diagram_base("Quyền được tạo ra như thế nào", "RBAC chỉ trả lời ai có thể thử làm gì ở phạm vi nào")
    base.box(d, (80, 260, 500, 470), "Principal", "Người dùng hoặc Group nhận quyền", fill="#EAF3FA", title_size=37, body_size=25)
    base.box(d, (580, 260, 1000, 470), "Role Definition", "Tập Permission được định nghĩa và có phiên bản", fill="#EAF3FA", title_size=35, body_size=25)
    base.box(d, (1080, 260, 1500, 470), "Scope", "Công ty, Project hoặc tài nguyên", fill="#EAF3FA", title_size=37, body_size=25)
    _label(d, (540, 365), "+", size=48, color="#526171", bold=True, anchor="mm")
    _label(d, (1040, 365), "+", size=48, color="#526171", bold=True, anchor="mm")
    base.box(d, (430, 590, 1150, 800), "Role Assignment", "Gắn đúng Principal, Role Definition và Scope; có trạng thái và lịch sử", fill="#EAF3FA", title_size=39, body_size=26)
    for x in (290, 790, 1290):
        base.arrow(d, (x, 470), (790, 590), color="#526171", width=4, head=15)
    base.arrow(d, (1150, 695), (1390, 695), color="#526171")
    base.box(d, (1390, 575, 1900, 815), "Kết quả", "Effective Permission\n+ điều kiện nghiệp vụ\n= cho phép hoặc từ chối", fill="#E8F5EE", outline="#2E7D5B", title_size=36, body_size=27)
    _label(d, (1000, 980), "Có Permission chưa đủ: tài liệu vẫn phải đúng trạng thái, đúng Checkout và đúng Generation.", size=25, color="#0F2747", bold=True, anchor="mm")
    figs["rbac"] = _save(img, "figure-10-rbac-model.png")

    # 11. Project-scoped example
    img, d = base.diagram_base("Ví dụ cấp quyền cho Linh trong Project P-100", "Tài khoản, tư cách thành viên và quyền làm việc là ba thông tin riêng")
    steps = [
        ("1. Tài khoản", "QLHT tạo và kích hoạt tài khoản Linh", "#F6F8FA", "#7B8FA3"),
        ("2. Project", "Project Administrator thêm Linh vào P-100", "#EAF3FA", "#1F6FB2"),
        ("3. Group", "Linh là thành viên trực tiếp của Group Cơ khí P-100", "#EAF3FA", "#1F6FB2"),
        ("4. Role Assignment", "Group Cơ khí nhận Role Design Engineer tại Scope P-100", "#EAF3FA", "#1F6FB2"),
        ("5. Khi thao tác", "Server tính Effective Permission rồi kiểm tra trạng thái tài liệu", "#E8F5EE", "#2E7D5B"),
    ]
    for i, (title, subtitle, fill, outline) in enumerate(steps):
        x = 50 + i * 390
        base.box(d, (x, 360, x + 340, 690), title, subtitle, fill=fill, outline=outline, title_size=29, body_size=24)
        if i:
            base.arrow(d, (x - 50, 525), (x, 525), color=outline, width=4, head=15)
    _label(d, (1000, 840), "Cùng một người có thể là Design Engineer ở P-100 nhưng chỉ là người xem ở Project khác.", size=28, color="#0F2747", bold=True, anchor="mm")
    _label(d, (1000, 920), "Gán trực tiếp cho một người vẫn được hỗ trợ như ngoại lệ và phải hiển thị trong Audit.", size=25, color="#526171", anchor="mm")
    figs["project_permission"] = _save(img, "figure-11-project-permission.png")

    # 12. Admin separation
    img, d = base.diagram_base("Tách trách nhiệm quản trị", "Quyền quản trị hệ thống không tự tạo quyền xem, sửa, duyệt hoặc Release tài liệu")
    admins = [
        ("Account Administrator", "Tạo, khóa, khôi phục tài khoản và phiên đăng nhập"),
        ("Project Administrator", "Quản lý thành viên Project, Group và assignment trong phạm vi"),
        ("Privileged Role Administrator", "Quản lý Role Definition và quyền quản trị được ủy quyền"),
        ("Product Configuration Administrator", "Quản lý loại tài liệu, workflow và policy có phiên bản"),
        ("Audit Reader", "Tra cứu lịch sử trong Scope; không sửa bằng chứng"),
        ("Super Administrator", "Khởi tạo và phục hồi khẩn cấp; không dùng hằng ngày"),
    ]
    for i, (title, subtitle) in enumerate(admins):
        col = i % 3
        row = i // 3
        x = 80 + col * 640
        y = 245 + row * 350
        fill = "#FCECEC" if title == "Super Administrator" else "#F6F9FC"
        outline = "#A43B3B" if title == "Super Administrator" else "#1F6FB2"
        base.box(d, (x, y, x + 560, y + 250), title, subtitle, fill=fill, outline=outline, title_size=31, body_size=24)
    _label(d, (1000, 1000), "Một người có thể kiêm nhiều vai trò, nhưng hệ thống vẫn ghi thành các Role Assignment riêng theo đúng Scope.", size=25, color="#0F2747", bold=True, anchor="mm")
    figs["admin"] = _save(img, "figure-12-admin-separation.png")

    # 13. Authorization decision
    img, d = base.diagram_base("Một yêu cầu được cho phép hay từ chối như thế nào", "Mỗi lần thao tác đều kiểm tra lại quyền hiện hành và điều kiện của tài liệu")
    xs = _sequence_lanes(d, ["Người dùng", "Identity và Accounts", "Access Policy", "Module tài liệu", "Audit"])
    y = 315
    _message(d, xs[0], xs[1], y, "Gửi yêu cầu với session")
    y += 100
    _message(d, xs[1], xs[2], y, "Actor hợp lệ, Project và Group hiện hành")
    y += 100
    _message(d, xs[2], xs[2] + 1, y, "Tìm Role Assignment áp dụng theo Scope", color="#A96C12")
    y += 100
    _message(d, xs[2], xs[3], y, "Effective Permission cho phép thử thao tác", color="#2E7D5B")
    y += 100
    _message(d, xs[3], xs[3] + 1, y, "Kiểm tra trạng thái, Checkout, Generation và policy", color="#A96C12")
    y += 100
    _message(d, xs[3], xs[4], y, "Ghi quyết định và kết quả")
    y += 100
    _message(d, xs[3], xs[0], y, "Cho phép hoặc từ chối có lý do", color="#2E7D5B", dashed=True)
    _label(d, (1000, 1040), "Nếu thiếu một điều kiện, hệ thống từ chối an toàn; không bỏ qua bước kiểm tra vì người dùng là quản trị viên.", size=23, color="#A43B3B", bold=True, anchor="mm")
    figs["auth_decision"] = _save(img, "figure-13-auth-decision.png")

    # 14. Operations without technology proposal
    img, d = base.diagram_base("Triển khai và bảo vệ dữ liệu ở mức logic", "Các vùng trách nhiệm cần có; báo cáo không chọn hệ điều hành, framework hoặc phiên bản sản phẩm")
    base.box(d, (70, 300, 470, 590), "Máy kỹ sư Windows", "Web, Desktop, Workspace và Office/CAD. Không giữ credential database hoặc kho file.", fill="#F6F9FC", title_size=33, body_size=24)
    base.box(d, (690, 260, 1310, 700), "Vùng dịch vụ do công ty quản lý", "IDEA Server; PostgreSQL; kho file; bộ xử lý định dạng có giới hạn.", fill="#EAF3FA", title_size=36, body_size=27)
    base.box(d, (1530, 250, 1930, 470), "Vùng backup độc lập", "Database, file, cấu hình, policy và key cùng một recovery set", fill="#E8F5EE", outline="#2E7D5B", title_size=31, body_size=24)
    base.box(d, (1530, 610, 1930, 850), "Mở rộng khi có số đo", "Tăng dung lượng, worker hoặc host khi tải, thời gian phục hồi hay vận hành yêu cầu", fill="#FFF4DB", outline="#A96C12", title_size=31, body_size=24)
    base.arrow(d, (470, 445), (690, 445), label="HTTPS")
    base.arrow(d, (1310, 390), (1530, 360), color="#2E7D5B", label="backup phối hợp")
    base.arrow(d, (1310, 600), (1530, 730), color="#A96C12", label="số đo")
    _label(d, (1000, 930), "Khôi phục chỉ đạt khi mở được đúng dữ liệu, đối soát đủ digest và không làm sống lại session hoặc quyền đã thu hồi.", size=24, color="#0F2747", bold=True, anchor="mm")
    figs["operation"] = _save(img, "figure-14-operation.png")

    return figs


def add_figure(doc: Document, path: Path, title: str, description: str, width_cm=18.0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    inline = p.add_run().add_picture(str(path), width=Cm(width_cm))
    base.set_inline_alt(inline, title, description)


def add_source_links(doc: Document, entries: list[tuple[str, Path]]):
    p = doc.add_paragraph(style="Small Note")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Mở sơ đồ kỹ thuật chi tiết: ")
    for idx, (label, path) in enumerate(entries):
        if idx:
            p.add_run("; ")
        base.add_hyperlink(p, label, path.resolve().as_uri())


def add_plain_note(doc: Document, lead: str, text: str):
    p = doc.add_paragraph()
    r = p.add_run(lead + ". ")
    r.bold = True
    p.add_run(text)
    return p


def _set_update_fields(doc: Document):
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def build_document(figs: dict[str, Path]) -> None:
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    if not LOGO.exists():
        raise FileNotFoundError(LOGO)
    HUMAN.mkdir(parents=True, exist_ok=True)
    WORK.mkdir(parents=True, exist_ok=True)

    logo_embed = WORK / "logo-idea-embed.png"
    with Image.open(LOGO) as logo_image:
        logo_image.save(logo_embed, format="PNG", optimize=True)

    doc = Document(str(REFERENCE))
    base.clear_body(doc)
    base.set_doc_defaults(doc)

    # Word's default TOC styles are too loose for this report and can leave a
    # nearly blank continuation page. Keep the same hierarchy but make it
    # compact enough for one page.
    for style_name, font_size, line_height in (("TOC 1", 9.2, 12.0), ("TOC 2", 8.8, 11.5), ("TOC 3", 8.5, 11.0)):
        try:
            toc_style = doc.styles[style_name]
        except KeyError:
            # The reference document has no populated TOC yet. Defining the
            # standard style names here lets Word reuse these settings when it
            # expands the field.
            toc_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        toc_style.font.name = "Calibri"
        toc_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        toc_style.font.size = Pt(font_size)
        toc_style.paragraph_format.space_before = Pt(0)
        toc_style.paragraph_format.space_after = Pt(0)
        toc_style.paragraph_format.line_spacing = Pt(line_height)
    if "Figure Caption" in doc.styles:
        doc.styles["Figure Caption"].paragraph_format.keep_with_next = True
    doc.core_properties.title = "Báo cáo kiến trúc nghiệp vụ IDEA DDM Core v0"
    doc.core_properties.subject = "Tài liệu trình review về phạm vi, khối hệ thống, kiểm soát tài liệu, RBAC và vận hành"
    doc.core_properties.keywords = "IDEA DDM, Checkout, Check-in, Reference, Review, Release, RBAC, Project, Permission"
    doc.core_properties.comments = "Dự thảo trình review kiến trúc nghiệp vụ; không đề xuất lựa chọn công nghệ."
    base.configure_cover_header_footer(doc.sections[0])

    # Cover
    cover = doc.add_table(rows=1, cols=2)
    cover.alignment = WD_TABLE_ALIGNMENT.CENTER
    cover.autofit = False
    cover.columns[0].width = Cm(12.5)
    cover.columns[1].width = Cm(5.0)
    base.remove_table_borders(cover)
    left, right = cover.rows[0].cells
    for label, value in (
        ("Người phụ trách", "Nguyễn Huỳnh Phúc Lâm"),
        ("Phòng ban", "Phòng Phát triển sản phẩm nội bộ"),
        ("Ngày soạn", "11/09/2026"),
        ("Trạng thái", "Dự thảo trình review nghiệp vụ và kiến trúc"),
    ):
        p = left.add_paragraph() if left.paragraphs[0].text else left.paragraphs[0]
        r = p.add_run(f"{label}: ")
        r.bold = True
        p.add_run(value)
        p.paragraph_format.space_after = Pt(3)
        base.set_paragraph_text_style(p, size=10.5, color=base.BLACK)
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo_inline = rp.add_run().add_picture(str(logo_embed), width=Cm(3.2))
    base.set_inline_alt(logo_inline, "Logo IDEA", "Logo nhận diện thương hiệu IDEA.")

    for _ in range(5):
        doc.add_paragraph().paragraph_format.space_after = Pt(7)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("BÁO CÁO KIẾN TRÚC NGHIỆP VỤ\nIDEA DDM CORE v0")
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Trọng tâm kiểm soát tài liệu và phân quyền")
    base.set_run_font(r, size=14, color=base.GRAY)
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p3.add_run("Tài liệu trình Product Decision Authority review")
    base.set_run_font(r, size=11, color=base.BLUE, bold=True)
    p3.paragraph_format.space_before = Pt(24)
    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    base.set_run_font(p.add_run("TÀI LIỆU NỘI BỘ"), size=9.5, color=base.GRAY, bold=True)

    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    base.set_section_layout(body_section)
    base.configure_body_header_footer(body_section)

    # Front matter
    doc.add_heading("Kiểm soát tài liệu", level=1)
    base.add_body_paragraph(doc, "Báo cáo này giúp người nắm nghiệp vụ xem xét cách IDEA DDM quản lý tài liệu, quyền sửa, xét duyệt, phát hành và quyền truy cập. Báo cáo không đề xuất framework, ngôn ngữ lập trình, hệ điều hành, phiên bản sản phẩm hoặc nhà cung cấp hạ tầng.")
    base.add_table_caption(doc, 1, "Thông tin kiểm soát báo cáo")
    base.add_table(doc, ["Thông tin", "Giá trị"], [
        ("Mã tài liệu", "IE-MGMT-ARCH-REPORT-001"),
        ("Phiên bản", "0.2 - Dự thảo trình review"),
        ("Phạm vi", "IDEA DDM Core v0"),
        ("Nguồn nghiệp vụ", "DOC-03@0.7 và DOC-04@0.13"),
        ("Nguồn kiến trúc", "DOC-05@0.13 và DOC-06@0.13"),
        ("Trọng tâm", "Checkout, Check-in, Reference; RBAC và trách nhiệm quản trị"),
        ("Không thuộc tài liệu", "Lựa chọn stack, phiên bản công nghệ, cấu hình production và tuyên bố kết quả kiểm thử"),
        ("Người lập và self-review", "Nguyễn Huỳnh Phúc Lâm"),
        ("Người duyệt", "Product Decision Authority - chưa ghi nhận quyết định"),
    ], widths=[4.2, 13.7], font_size=9.8)

    doc.add_heading("Nội dung đề nghị review", level=2)
    base.add_table_caption(doc, 2, "Năm nội dung cần xem xét")
    base.add_table(doc, ["Phần", "Câu hỏi cần trả lời", "Kết quả mong muốn"], [
        ("I", "Phần mềm giải quyết công việc gì và ai sử dụng?", "Xác nhận đúng vấn đề, nhóm người dùng và ranh giới sản phẩm."),
        ("II", "Sáu khối chính có chia trách nhiệm rõ không?", "Xác nhận Web, Desktop, Server, PostgreSQL, kho file và bộ xử lý định dạng không tự làm thay quyền của nhau."),
        ("III", "Checkout, Check-in, Reference, Review và Release có đúng nghiệp vụ không?", "Xác nhận phạm vi, kết quả thành công, cách từ chối và cách giữ an toàn dữ liệu khi có lỗi."),
        ("IV", "RBAC và trách nhiệm quản trị có đúng cách tổ chức công việc không?", "Xác nhận quyền theo Project, Group, Role và Scope; quản trị tài khoản không đồng nghĩa có quyền tài liệu."),
        ("V", "Triển khai, backup, bảo mật và mở rộng đã có ranh giới hợp lý chưa?", "Xác nhận trách nhiệm và tiêu chí cần kiểm chứng; không chọn công nghệ trong báo cáo này."),
    ], widths=[1.2, 8.2, 8.5], font_size=9.6)
    add_plain_note(doc, "Kết luận hiện tại", "Thiết kế đủ để review ở mức nghiệp vụ và kiến trúc logic. Chưa có bằng chứng để tuyên bố hệ thống đã triển khai, đạt hiệu năng, phục hồi thành công hoặc an toàn cho production.")
    doc.add_page_break()

    doc.add_heading("Mục lục", level=1)
    p = doc.add_paragraph()
    base.add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Mục lục được cập nhật khi mở tài liệu trong Microsoft Word.")
    doc.add_page_break()
    doc.add_heading("Danh mục hình", level=1)
    p = doc.add_paragraph()
    base.add_field(p, 'TOC \\h \\z \\t "Figure Caption,1"', "Danh mục hình được cập nhật khi mở tài liệu trong Microsoft Word.")
    doc.add_page_break()
    doc.add_heading("Danh mục bảng", level=1)
    p = doc.add_paragraph()
    base.add_field(p, 'TOC \\h \\z \\t "Table Caption,1"', "Danh mục bảng được cập nhật khi mở tài liệu trong Microsoft Word.")
    doc.add_page_break()

    # Part I
    doc.add_heading("Phần I Phần mềm giải quyết công việc gì và ai sử dụng", level=1)
    base.add_body_paragraph(doc, "IDEA DDM là phần mềm nội bộ quản lý tài liệu và dữ liệu thiết kế có kiểm soát. Kỹ sư vẫn soạn bằng Office hoặc CAD. IDEA DDM quản lý danh tính tài liệu, bản nội dung, cấu trúc sản phẩm, người được phép thao tác, quá trình review và bộ hồ sơ được Release.")
    add_figure(doc, figs["users"], "Người dùng IDEA DDM", FIGURE_ALT["figure-01-users.png"])
    base.add_figure_caption(doc, 1, "Người dùng và công việc chính trong IDEA DDM")
    add_source_links(doc, [("ARCH-VIEW-CTX-001", EVIDENCE / "ARCH-VIEW-CTX-001.svg")])
    doc.add_heading("1.1 Người sử dụng", level=2)
    base.add_table_caption(doc, 3, "Người sử dụng và trách nhiệm chính")
    base.add_table(doc, ["Người sử dụng", "Công việc trong IDEA DDM", "Điều không được hiểu nhầm"], [
        ("Kỹ sư hoặc người soạn", "Tìm tài liệu, lấy file bằng Checkout hoặc Reference, Check-in và gửi review.", "Quyền sửa một tài liệu không tự mở rộng sang cả cây sản phẩm."),
        ("Người review", "Xem đúng Generation được gửi; Approve hoặc Reject có lý do.", "Không review một bản có thể âm thầm đổi nội dung sau quyết định."),
        ("Người có quyền Release", "Xem và xác nhận đúng phạm vi; hệ thống kiểm tra lại điều kiện trước khi phát hành.", "Under Review chưa đồng nghĩa được phép Release."),
        ("Quản trị tài khoản", "Tạo, khóa và khôi phục tài khoản IDEA.", "Không vì quản trị tài khoản mà có quyền xem hoặc sửa tài liệu."),
        ("Quản trị Project, Role hoặc cấu hình", "Thực hiện phần quản trị được giao tại đúng Scope.", "Mỗi trách nhiệm là một Role Assignment riêng, có Audit."),
        ("Người tra cứu", "Mở bản được phép dùng, xem lịch sử và bằng chứng.", "Audit không phải quyền sửa lịch sử."),
    ], widths=[3.6, 7.8, 6.5], font_size=9.3)

    doc.add_heading("1.2 Công việc hoàn chỉnh", level=2)
    base.add_numbered(doc, [
        "Đưa tài liệu hiện có vào hệ thống hoặc tạo tài liệu mới với một mã ổn định.",
        "Lấy đúng file về Workspace bằng Checkout để sửa hoặc Reference để tham khảo.",
        "Check-in thay đổi theo một phạm vi đã xác nhận và không để lại kết quả thành công một phần.",
        "Gửi đúng Generation để review, ghi người thực hiện, thời điểm và quyết định.",
        "Release đúng bộ hồ sơ đã xác nhận và có thể lấy lại đúng bộ đó về sau.",
    ])
    add_plain_note(doc, "Phạm vi Core v0", "Sản phẩm tập trung vào quản lý tài liệu và cấu trúc kỹ thuật. Phần mềm chưa thay Office/CAD và chưa tự điều khiển quá trình thiết kế hoặc nghiệp vụ ERP.")

    # Part II
    doc.add_heading("Phần II Các khối chính của hệ thống", level=1)
    base.add_body_paragraph(doc, "Sáu khối được tách theo trách nhiệm. Web và Desktop nhận thao tác. Server là nơi duy nhất ra quyết định nghiệp vụ. PostgreSQL lưu dữ liệu giao dịch; kho file lưu nội dung; bộ xử lý định dạng tạo bản xem trước hoặc bản dẫn xuất trong vùng được giới hạn.")
    add_figure(doc, figs["blocks"], "Sáu khối chính", FIGURE_ALT["figure-02-blocks.png"])
    base.add_figure_caption(doc, 2, "Sáu khối chính và ranh giới trách nhiệm")
    add_source_links(doc, [("ARCH-VIEW-CON-001", EVIDENCE / "ARCH-VIEW-CON-001.svg"), ("ARCH-VIEW-MOD-001", EVIDENCE / "ARCH-VIEW-MOD-001.svg")])
    doc.add_heading("2.1 Trách nhiệm từng khối", level=2)
    base.add_table_caption(doc, 4, "Trách nhiệm và giới hạn của sáu khối")
    base.add_table(doc, ["Khối", "Làm gì", "Không được tự làm gì"], [
        ("Web", "Tìm kiếm, xem, review, Release và mở màn hình quản trị được cấp quyền.", "Không tự tính quyền hoặc ghi trực tiếp vào PostgreSQL và kho file."),
        ("Desktop", "Quản lý Workspace, tải file, mở Office/CAD và gửi thao tác có xác nhận.", "Không tự công bố file khi người dùng chỉ bấm Save trong ứng dụng soạn thảo."),
        ("IDEA Server", "Xác thực, RBAC, Checkout, Check-in, workflow, Release và Audit.", "Không giao quyền quyết định cho client, database script hoặc bộ xử lý định dạng."),
        ("PostgreSQL", "Lưu metadata, trạng thái, quan hệ, quyền, workflow, Release Record và Audit theo giao dịch.", "Không phải nơi người dùng sửa file làm việc và không tự quyết định nghiệp vụ."),
        ("Kho file", "Lưu nội dung bất biến và trả file theo yêu cầu đã được Server kiểm tra.", "Không phải ổ dùng chung để người dùng sửa trực tiếp."),
        ("Bộ xử lý định dạng", "Đọc file bất biến để tạo preview hoặc Representation theo job có giới hạn.", "Không sửa file nguồn và không đổi trạng thái tài liệu."),
    ], widths=[3.1, 7.7, 7.1], font_size=9.3)
    add_plain_note(doc, "Giới hạn của phần này", "Báo cáo chỉ mô tả khối và thẩm quyền. Việc chọn ngôn ngữ lập trình, framework, hệ điều hành, phiên bản PostgreSQL, loại kho file hoặc công cụ chuyển đổi nằm ngoài phạm vi review này.")

    # Part III
    doc.add_heading("Phần III Luồng tài liệu hoàn chỉnh", level=1)
    base.add_body_paragraph(doc, "Phần này tách ba khái niệm thường bị nhầm. Checkout là quyền giữ sửa. Reference là bản tham khảo. Check-in là thao tác có xác nhận để công bố thay đổi từ Workspace vào hệ thống. Sau đó một Generation cụ thể mới được gửi Review và có thể được Release khi đủ điều kiện.")

    doc.add_heading("3.1 Chọn Checkout hoặc Reference", level=2)
    add_figure(doc, figs["choice"], "Chọn Checkout hoặc Reference", FIGURE_ALT["figure-03-checkout-reference-choice.png"])
    base.add_figure_caption(doc, 3, "Cách chọn Checkout hoặc Reference")
    add_source_links(doc, [("ARCH-VIEW-SEQ-001", EVIDENCE / "ARCH-VIEW-SEQ-001.svg")])
    base.add_table_caption(doc, 5, "Phân biệt Checkout và Reference")
    base.add_table(doc, ["Nội dung", "Checkout", "Reference"], [
        ("Mục đích", "Sửa tài liệu rồi Check-in.", "Xem, đối chiếu hoặc dùng làm đầu vào tham khảo."),
        ("Quyền giữ sửa", "Có, cho đúng người và đúng Workspace trong thời hạn.", "Không."),
        ("Người khác", "Không thể Checkout cùng tài liệu khi quyền giữ còn hiệu lực.", "Vẫn có thể lấy Reference."),
        ("Check-in vào tài liệu gốc", "Có, nếu mọi điều kiện vẫn hợp lệ.", "Không; phải xin Checkout hợp lệ trước."),
        ("Phạm vi", "Chỉ tài liệu được người dùng xác nhận; không tự giữ cả cây.", "Chỉ những file được chọn để tham khảo."),
    ], widths=[4.1, 6.9, 6.9], font_size=9.4)

    doc.add_heading("3.2 Trình tự Checkout", level=2)
    add_figure(doc, figs["checkout"], "Trình tự Checkout", FIGURE_ALT["figure-04-checkout-sequence.png"])
    base.add_figure_caption(doc, 4, "Trình tự Checkout đã được xác nhận")
    add_source_links(doc, [("ARCH-VIEW-SEQ-001", EVIDENCE / "ARCH-VIEW-SEQ-001.svg"), ("ARCH-VIEW-STATE-002", EVIDENCE / "ARCH-VIEW-STATE-002.svg")])
    base.add_body_paragraph(doc, "Checkout chỉ hoàn tất khi Server đã xác nhận tài khoản, quyền tại đúng Scope, trạng thái tài liệu, expected Generation và việc chưa có Workspace khác giữ sửa. Sau đó Workspace mới tải file, kiểm tra digest và ghi manifest cục bộ.")

    doc.add_heading("3.3 Reference và file Reference đã bị sửa", level=2)
    add_figure(doc, figs["reference"], "Xử lý Reference", FIGURE_ALT["figure-05-reference-handling.png"])
    base.add_figure_caption(doc, 5, "Cách xử lý an toàn một file Reference")
    add_source_links(doc, [("ARCH-VIEW-STATE-003", EVIDENCE / "ARCH-VIEW-STATE-003.svg"), ("ARCH-VIEW-SEQ-005", EVIDENCE / "ARCH-VIEW-SEQ-005.svg")])
    base.add_body_paragraph(doc, "Nếu người dùng sửa một file Reference, IDEA DDM không coi đó là thay đổi chính thức. Khi bản nguồn vẫn hiện hành và chưa bị người khác Checkout, người dùng có thể xin Checkout đúng Generation rồi tiếp tục với file cục bộ. Nếu bản nguồn đã mới hơn, hệ thống giữ file để người dùng so sánh và áp dụng lại thay đổi có chủ đích; không tự merge CAD hoặc Office.")

    doc.add_heading("3.4 Check-in và các kết quả", level=2)
    add_figure(doc, figs["checkin"], "Kết quả Check-in", FIGURE_ALT["figure-06-checkin-outcomes.png"])
    base.add_figure_caption(doc, 6, "Bốn kết quả cần phân biệt khi Check-in")
    add_source_links(doc, [("ARCH-VIEW-SEQ-002", EVIDENCE / "ARCH-VIEW-SEQ-002.svg"), ("ARCH-VIEW-STATE-004", EVIDENCE / "ARCH-VIEW-STATE-004.svg")])
    base.add_table_caption(doc, 6, "Kết quả Check-in và ảnh hưởng đến dữ liệu")
    base.add_table(doc, ["Kết quả", "Dữ liệu chính thức", "Checkout và file cục bộ"], [
        ("Có thay đổi", "Tạo một Generation mới cho mỗi tài liệu thay đổi; tăng Version trong Revision.", "Kết thúc Checkout của toàn bộ phạm vi đã xác nhận."),
        ("No Change", "Không tạo Generation và không tăng Version.", "Vẫn kết thúc Checkout của toàn bộ phạm vi đã xác nhận."),
        ("Từ chối trước commit", "Không tạo bản mới và không cập nhật một phần.", "Giữ file cục bộ; Checkout chỉ đổi nếu có thao tác được phép riêng."),
        ("Chưa xác định do gián đoạn", "Chưa được tự kết luận thành công hoặc thất bại.", "Giữ file cục bộ; tra cứu cùng OperationId để biết đã commit hay có thể tiếp tục."),
    ], widths=[3.4, 7.2, 7.3], font_size=9.3)
    add_plain_note(doc, "Quy tắc sau Check-in", "Check-in thành công, kể cả No Change, phải kết thúc quyền giữ sửa trong toàn bộ phạm vi người dùng đã xác nhận. Check-in nhiều tài liệu có một kết quả all-or-none.")

    doc.add_heading("3.5 Bản làm việc đã cũ và thao tác bị gián đoạn", level=2)
    add_figure(doc, figs["recovery"], "Xử lý bản cũ và gián đoạn", FIGURE_ALT["figure-07-stale-recovery.png"])
    base.add_figure_caption(doc, 7, "Xử lý bản làm việc đã cũ và Check-in bị gián đoạn")
    add_source_links(doc, [("ARCH-VIEW-SEQ-007", EVIDENCE / "ARCH-VIEW-SEQ-007.svg"), ("ARCH-VIEW-MOD-002", EVIDENCE / "ARCH-VIEW-MOD-002.svg")])
    base.add_table_caption(doc, 7, "Cách xử lý khi Check-in không đi theo luồng bình thường")
    base.add_table(doc, ["Tình huống", "Hệ thống phải làm", "Người dùng nhìn thấy"], [
        ("Expected Generation đã cũ", "Từ chối ghi; giữ tài liệu chính thức nguyên vẹn và giữ file cục bộ.", "Bản hiện tại, bản mình đã dựa vào và các lựa chọn so sánh hoặc áp dụng lại."),
        ("Mất kết nối trước commit", "Ghi nhận tiến độ đã xác minh theo OperationId; cho phép tiếp tục phần còn thiếu.", "Không phải tải lại phần đã được xác nhận."),
        ("Commit xong nhưng mất phản hồi", "Trả lại đúng kết quả đã commit khi client hỏi cùng OperationId.", "Không tạo Generation thứ hai."),
        ("Bằng chứng kết quả chưa nhất quán", "Giữ thao tác để đối soát; không cho gửi một thao tác mới gây trùng.", "Trạng thái cần kiểm tra, file cục bộ vẫn an toàn."),
    ], widths=[4.1, 7.2, 6.6], font_size=9.2)

    doc.add_heading("3.6 Review và Release", level=2)
    add_figure(doc, figs["lifecycle"], "Vòng đời tài liệu", FIGURE_ALT["figure-08-document-lifecycle.png"])
    base.add_figure_caption(doc, 8, "Vòng đời từ Start đến Released")
    add_source_links(doc, [("ARCH-VIEW-STATE-001", EVIDENCE / "ARCH-VIEW-STATE-001.svg"), ("ARCH-VIEW-SEQ-003", EVIDENCE / "ARCH-VIEW-SEQ-003.svg")])
    base.add_table_caption(doc, 8, "Ý nghĩa Review và Release")
    base.add_table(doc, ["Bước", "Ghi nhận gì", "Điều kiện quan trọng"], [
        ("Submit for Review", "Gửi một Generation và phạm vi cụ thể vào Review Round.", "Bản gửi phải được pin; thay đổi sau đó không kế thừa quyết định cũ."),
        ("Approve hoặc Reject", "Ghi người quyết định, thời điểm, bản được xem và lý do khi Reject.", "Người review phải đủ điều kiện và độc lập theo policy."),
        ("Release", "Tạo Release Record và package bất biến cho đúng phạm vi.", "Kiểm tra lại quyền, trạng thái, phụ thuộc bắt buộc và bằng chứng ngay lúc commit."),
    ], widths=[4.1, 7.0, 6.8], font_size=9.3)

    doc.add_heading("3.7 Phát hành từng phần trong P-100", level=2)
    add_figure(doc, figs["staged_release"], "Phát hành theo phạm vi P-100", FIGURE_ALT["figure-09-staged-release.png"])
    base.add_figure_caption(doc, 9, "Cụm bơm được Release trong khi các phần khác vẫn In Work")
    add_source_links(doc, [("ARCH-VIEW-ACT-002", EVIDENCE / "ARCH-VIEW-ACT-002.svg"), ("ARCH-VIEW-SEQ-003", EVIDENCE / "ARCH-VIEW-SEQ-003.svg")])
    base.add_table_caption(doc, 9, "Ví dụ phạm vi Release của P-100")
    base.add_table(doc, ["Phạm vi", "Trạng thái có thể có", "Cách hệ thống xử lý"], [
        ("Cụm bơm", "Released", "Release Record pin đúng Generation, Structure Snapshot, phụ thuộc bắt buộc và bằng chứng của cụm bơm."),
        ("Tủ điện", "In Work", "Tiếp tục thiết kế; không bị kéo sang Released chỉ vì cùng thuộc P-100."),
        ("Hồ sơ toàn máy", "In Work", "Chỉ được Release khi phạm vi riêng của nó đủ điều kiện."),
    ], widths=[3.7, 4.2, 10.0], font_size=9.4)

    # Part IV
    doc.add_heading("Phần IV RBAC và trách nhiệm quản trị", level=1)
    part_iv_intro = base.add_body_paragraph(doc, "RBAC cấp quyền bằng Role Assignment. Một Role Assignment gắn Principal với một Role Definition tại một Scope. Principal có thể là người dùng hoặc Group. Khi người dùng thao tác, hệ thống tổng hợp các assignment còn hiệu lực rồi kiểm tra thêm điều kiện nghiệp vụ của tài liệu.")
    part_iv_intro.paragraph_format.keep_with_next = True

    doc.add_heading("4.1 Mô hình RBAC", level=2)
    add_figure(doc, figs["rbac"], "Mô hình RBAC", FIGURE_ALT["figure-10-rbac-model.png"])
    base.add_figure_caption(doc, 10, "Principal cộng Role Definition cộng Scope tạo Role Assignment")
    add_source_links(doc, [("ARCH-VIEW-RBAC-001", EVIDENCE / "ARCH-VIEW-RBAC-001.svg"), ("DATA-VIEW-AUTH-001", EVIDENCE / "DATA-VIEW-AUTH-001.svg")])
    base.add_table_caption(doc, 10, "Các thành phần của RBAC")
    base.add_table(doc, ["Khái niệm", "Cách hiểu", "Ví dụ P-100"], [
        ("Security Principal", "Đối tượng có thể nhận Role Assignment: một người dùng hoặc Group.", "Linh Nguyễn hoặc Group Cơ khí P-100."),
        ("Role Definition", "Bộ Permission có phiên bản.", "Design Engineer có quyền xem, Checkout và Check-in theo policy."),
        ("Permission", "Một hành động mà Server hiểu và kiểm tra.", "Document.Checkout hoặc Document.CheckIn."),
        ("Authorization Scope", "Phạm vi Role có hiệu lực.", "Project P-100 hoặc một tài nguyên cụ thể."),
        ("Role Assignment", "Bản ghi gắn Principal, Role Definition và Scope.", "Group Cơ khí nhận Role Design Engineer tại P-100."),
        ("Effective Permission", "Quyền hợp lệ tại thời điểm yêu cầu sau khi xét assignment trực tiếp và qua Group.", "Linh có thể thử Checkout trong P-100."),
        ("Business gate", "Điều kiện của tài liệu vẫn phải đạt sau RBAC.", "Tài liệu chưa Released, chưa bị Workspace khác giữ và đúng expected Generation."),
    ], widths=[3.7, 7.4, 6.8], font_size=9.2)

    doc.add_heading("4.2 Quyền thay đổi theo Project", level=2)
    add_figure(doc, figs["project_permission"], "Quyền theo Project", FIGURE_ALT["figure-11-project-permission.png"])
    base.add_figure_caption(doc, 11, "Ví dụ cấp quyền cho Linh trong Project P-100")
    add_source_links(doc, [("ARCH-VIEW-ACT-001", EVIDENCE / "ARCH-VIEW-ACT-001.svg")])
    base.add_table_caption(doc, 11, "Các bước cấp quyền thông thường")
    base.add_table(doc, ["Bước", "Người thực hiện", "Kết quả"], [
        ("Tạo tài khoản", "Account Administrator thuộc QLHT", "Linh có thể đăng nhập nhưng chưa mặc nhiên có quyền vào P-100."),
        ("Thêm vào Project", "Project Administrator", "Linh trở thành thành viên P-100."),
        ("Thêm vào Group", "Project Administrator trong Scope P-100", "Linh là thành viên trực tiếp của Group Cơ khí P-100."),
        ("Gán Role", "Người được phép quản lý assignment", "Group Cơ khí nhận Role Design Engineer tại Scope P-100."),
        ("Thao tác tài liệu", "IDEA Server", "Tính Effective Permission rồi kiểm tra business gate của tài liệu."),
    ], widths=[3.2, 5.1, 9.6], font_size=9.3)
    add_plain_note(doc, "Hệ quả", "Linh có thể có vai trò khác nhau ở các Project khác nhau. Hệ thống không suy quyền từ chức danh, phòng ban hoặc việc có tài khoản.")

    doc.add_heading("4.3 Tách trách nhiệm quản trị", level=2)
    add_figure(doc, figs["admin"], "Trách nhiệm quản trị", FIGURE_ALT["figure-12-admin-separation.png"])
    base.add_figure_caption(doc, 12, "Các trách nhiệm quản trị được tách theo Role Assignment")
    add_source_links(doc, [("ARCH-VIEW-ACT-001", EVIDENCE / "ARCH-VIEW-ACT-001.svg"), ("ARCH-VIEW-STATE-005", EVIDENCE / "ARCH-VIEW-STATE-005.svg")])
    base.add_table_caption(doc, 12, "Vai trò quản trị và giới hạn")
    base.add_table(doc, ["Vai trò quản trị", "Được làm", "Không tự được làm"], [
        ("Account Administrator", "Tạo, kích hoạt, khóa, khôi phục tài khoản và phiên đăng nhập.", "Không thêm người vào Project hoặc tự cấp quyền xem, sửa, duyệt, Release."),
        ("Project Administrator", "Quản lý thành viên Project, Group và assignment được phép trong Project.", "Không quản trị tài khoản và không gán Role ngoài danh mục hoặc Scope được giao."),
        ("Privileged Role Administrator", "Quản lý Role Definition và Role Assignment quản trị trong giới hạn được giao.", "Không tự mở rộng quyền hoặc bỏ đường phục hồi cuối cùng."),
        ("Product Configuration Administrator", "Quản lý loại tài liệu, workflow và policy có phiên bản.", "Không cấp tài khoản, Project access hoặc quyền duyệt và Release theo mặc định."),
        ("Audit Reader", "Đọc Audit trong Scope được cấp.", "Không sửa hoặc xóa Audit Evidence."),
        ("Super Administrator", "Khởi tạo và phục hồi quyền quản trị cao nhất khi cần.", "Không dùng như tài khoản làm việc hằng ngày và không mặc nhiên có quyền nội dung."),
    ], widths=[4.2, 7.0, 6.7], font_size=9.0)

    doc.add_heading("4.4 Kiểm tra quyền tại thời điểm thao tác", level=2)
    add_figure(doc, figs["auth_decision"], "Kiểm tra quyền thao tác", FIGURE_ALT["figure-13-auth-decision.png"])
    base.add_figure_caption(doc, 13, "Trình tự tính quyền và kiểm tra điều kiện nghiệp vụ")
    add_source_links(doc, [("ARCH-VIEW-SEQ-004", EVIDENCE / "ARCH-VIEW-SEQ-004.svg")])
    base.add_table_caption(doc, 13, "Các lớp kiểm tra trước khi cho phép thao tác")
    base.add_table(doc, ["Lớp kiểm tra", "Câu hỏi", "Ví dụ từ chối"], [
        ("Tài khoản và session", "Người dùng có còn được đăng nhập không?", "Tài khoản đã khóa hoặc session đã bị thu hồi."),
        ("Project và Group", "Người dùng có còn là thành viên đúng Project hoặc Group không?", "Linh đã được rút khỏi Group Cơ khí P-100."),
        ("Role Assignment", "Có Role Definition phù hợp tại đúng Scope và còn hiệu lực không?", "Role chỉ có hiệu lực ở Project khác hoặc đã hết hạn."),
        ("Permission", "Role có chứa đúng hành động cần thực hiện không?", "Có quyền xem nhưng không có Document.CheckIn."),
        ("Business gate", "Trạng thái tài liệu và điều kiện nghiệp vụ có cho phép không?", "Tài liệu Released, Checkout thuộc Workspace khác hoặc expected Generation đã cũ."),
        ("Audit", "Quyết định và kết quả đã được ghi đủ chưa?", "Thiếu bằng chứng bắt buộc thì thao tác bảo vệ không được coi là hoàn tất."),
    ], widths=[3.8, 7.3, 6.8], font_size=9.2)

    # Part V
    doc.add_heading("Phần V Triển khai, backup, bảo mật và khả năng mở rộng", level=1)
    part_v_intro = base.add_body_paragraph(doc, "Phần này mô tả các vùng trách nhiệm và điều kiện cần kiểm chứng. Báo cáo không chọn hệ điều hành, framework, cách đóng gói, loại máy chủ, sản phẩm backup hoặc công cụ giám sát.")
    part_v_intro.paragraph_format.keep_with_next = True
    add_figure(doc, figs["operation"], "Triển khai logic và bảo vệ dữ liệu", FIGURE_ALT["figure-14-operation.png"])
    base.add_figure_caption(doc, 14, "Triển khai logic, backup độc lập, bảo mật và đường mở rộng")
    add_source_links(doc, [("ARCH-VIEW-DEP-001", EVIDENCE / "ARCH-VIEW-DEP-001.svg"), ("ARCH-VIEW-SEQ-011", EVIDENCE / "ARCH-VIEW-SEQ-011.svg"), ("ARCH-VIEW-SEC-001", EVIDENCE / "ARCH-VIEW-SEC-001.svg")])

    doc.add_heading("5.1 Ranh giới triển khai", level=2)
    base.add_table_caption(doc, 14, "Các vùng triển khai và trách nhiệm")
    base.add_table(doc, ["Vùng", "Thành phần", "Nguyên tắc"], [
        ("Máy kỹ sư Windows", "Web, Desktop, Workspace và Office/CAD.", "Không có credential PostgreSQL hoặc kho file; file làm việc nằm trong Workspace của người dùng."),
        ("Vùng dịch vụ công ty", "IDEA Server, PostgreSQL, kho file và job nền.", "Server là điểm quyết định; tách process identity và quyền file theo trách nhiệm."),
        ("Vùng xử lý định dạng", "Tiến trình hoặc máy chạy job đọc/chuyển đổi file.", "Nhận input bất biến, có giới hạn tài nguyên và không có quyền đổi product state."),
        ("Vùng backup độc lập", "Database, file, cấu hình, policy và key cần thiết.", "Không dùng bản sao cùng miền sự cố làm bằng chứng có thể phục hồi."),
    ], widths=[4.0, 6.5, 7.4], font_size=9.3)

    doc.add_heading("5.2 Backup và phục hồi", level=2)
    base.add_table_caption(doc, 15, "Một recovery set phải gồm những gì")
    base.add_table(doc, ["Nội dung", "Phải giữ", "Cách xác nhận phục hồi"], [
        ("PostgreSQL", "Dữ liệu chính thức và nhật ký cần để phục hồi đúng điểm.", "Mở được đúng trạng thái đã chọn và không bỏ qua lỗi nhật ký."),
        ("Kho file", "Mọi Artifact bất biến, manifest và digest.", "Mỗi Generation trong phạm vi đều tìm thấy file và digest khớp."),
        ("Cấu hình và policy", "Đúng phiên bản quyền, workflow và quy tắc đang hiệu lực tại điểm phục hồi.", "Không dùng sai policy hoặc không đọc được dữ liệu vì thiếu cấu hình."),
        ("Key và bí mật cần thiết", "Vật liệu được bảo vệ và quy trình truy cập.", "Khôi phục trong vùng cô lập; không để key trong mã nguồn hoặc bản backup mở."),
        ("Session và thay đổi bảo mật", "Vô hiệu session phục hồi và đối soát các quyền đã thu hồi sau điểm backup.", "Không làm sống lại tài khoản hoặc quyền đã bị khóa."),
        ("Restore drill", "Kịch bản phục hồi hoàn chỉnh và người chịu trách nhiệm.", "Đăng nhập, mở file, kiểm tra Release và đối soát digest trước khi mở dịch vụ."),
    ], widths=[3.5, 7.2, 7.2], font_size=9.2)

    doc.add_heading("5.3 Bảo mật theo ranh giới", level=2)
    base.add_table_caption(doc, 16, "Kiểm soát bảo mật cần có")
    base.add_table(doc, ["Ranh giới", "Kiểm soát", "Bằng chứng cần có trước khi dùng thật"], [
        ("Web hoặc Desktop đến Server", "Kênh mã hóa; session được bảo vệ; kiểm tra dữ liệu đầu vào; đánh giá lại tài khoản, RBAC và business gate trước khi ghi.", "Ma trận cho phép và từ chối qua từng bề mặt; thử session cũ, thao tác lặp và sai Scope."),
        ("Client đến PostgreSQL hoặc kho file", "Không cấp kết nối trực tiếp hay credential dài hạn; mọi truy cập đi qua Server.", "Thử truy cập trực tiếp bị từ chối; credential không xuất hiện trong client, log hoặc source."),
        ("Server đến bộ xử lý định dạng", "Input bất biến; giới hạn thời gian, bộ nhớ, kích thước; kết quả phải kiểm tra trước khi gắn với Generation.", "File lỗi, quá lớn hoặc job treo không làm đổi file nguồn và product state."),
        ("Quản trị và Audit", "Role Assignment có Scope; chống tự nâng quyền; Audit chỉ bổ sung cho thao tác bảo vệ.", "Thử vượt Scope, tự gán quyền, sửa Audit và bỏ đường phục hồi cuối cùng đều bị từ chối."),
        ("Runtime đến backup", "Quyền backup tách khỏi quyền runtime; phục hồi trong vùng cô lập.", "Restore drill hoàn chỉnh và danh sách người có quyền truy cập backup và key."),
    ], widths=[3.7, 8.2, 6.0], font_size=9.0)

    doc.add_heading("5.4 Khả năng mở rộng", level=2)
    base.add_table_caption(doc, 17, "Đường mở rộng và điều kiện kích hoạt")
    base.add_table(doc, ["Nhu cầu", "Thiết kế đã chừa đường", "Chỉ mở rộng khi"], [
        ("50 đến 100 người dùng tại một địa điểm", "Giao dịch dài được tách khỏi phản hồi tương tác; module có ranh giới rõ.", "Có số đo người dùng đồng thời, độ trễ, lỗi, khóa và tài nguyên."),
        ("Một file có thể nhiều GB", "Tải theo phần, kiểm tra digest, tiếp tục bằng OperationId; không buộc toàn bộ file nằm trong bộ nhớ.", "Có bộ file đại diện, ngưỡng thời gian và bộ nhớ, cùng thử nghiệm gián đoạn."),
        ("Một Project có thể hàng trăm TB", "Danh tính tài liệu không chứa đường dẫn vật lý; vị trí file được quản lý riêng.", "Dung lượng, thông lượng, phục hồi hoặc vận hành của kho hiện tại không còn đáp ứng."),
        ("Nhiều job xử lý định dạng", "Job không sở hữu trạng thái tài liệu và có thể chạy ở vùng tách biệt.", "Hàng chờ, thời gian xử lý, giấy phép hoặc nhu cầu cách ly cho thấy cần tăng worker."),
        ("Nhiều máy phục vụ", "Các module có interface và dữ liệu sở hữu rõ.", "Số đo tải hoặc nhu cầu cô lập chứng minh lợi ích; không tách chỉ vì dự kiến hệ thống lớn."),
    ], widths=[3.8, 8.1, 6.0], font_size=9.1)

    # Conclusion
    doc.add_heading("Kết luận và nội dung cần sếp xác nhận", level=1)
    base.add_body_paragraph(doc, "Báo cáo đã làm rõ năm nội dung ở mức nghiệp vụ và kiến trúc logic. Hai phần cần review kỹ nhất là cách giữ an toàn công việc khi Checkout, Reference và Check-in; và cách cấp quyền theo Project, Group, Role Definition và Scope.")
    base.add_table_caption(doc, 18, "Kết luận theo năm nội dung review")
    base.add_table(doc, ["Nội dung", "Kết luận hiện tại", "Điều cần sếp xác nhận"], [
        ("1 Công việc và người dùng", "IDEA DDM quản lý tài liệu và cấu trúc kỹ thuật từ lúc làm việc đến Release.", "Đúng vấn đề, đúng nhóm người dùng và đúng ranh giới sản phẩm."),
        ("2 Các khối chính", "Web và Desktop nhận thao tác; Server quyết định; PostgreSQL, kho file và bộ xử lý định dạng có trách nhiệm riêng.", "Cách chia trách nhiệm có phù hợp. Không yêu cầu chọn công nghệ trong báo cáo này."),
        ("3 Luồng tài liệu", "Checkout và Reference khác nhau; Check-in all-or-none; lỗi giữ file cục bộ; Review và Release pin đúng bản.", "Quy tắc thành công, từ chối, bản cũ, gián đoạn và Release từng phần đúng nghiệp vụ."),
        ("4 RBAC", "Principal cộng Role Definition cộng Scope tạo Role Assignment; quyền còn phải qua business gate.", "Mô hình theo Project và Group, cùng việc tách trách nhiệm quản trị, đúng cách tổ chức của công ty."),
        ("5 Vận hành", "Có vùng dịch vụ, vùng xử lý định dạng và vùng backup độc lập; mở rộng dựa trên số đo.", "Ranh giới trách nhiệm và bằng chứng cần có là phù hợp để tiếp tục thiết kế chi tiết."),
    ], widths=[3.5, 8.2, 6.2], font_size=9.2)

    # Standards
    doc.add_heading("Cơ sở tiêu chuẩn và cách đọc sơ đồ", level=1)
    base.add_body_paragraph(doc, "Tài liệu tổ chức nội dung theo ISO/IEC/IEEE 42010:2022: xác định người quan tâm, vấn đề cần xem xét, góc nhìn và sơ đồ trả lời từng câu hỏi. Yêu cầu nguồn được quản lý theo định hướng ISO/IEC/IEEE 29148:2018. Các đặc tính chất lượng được đối chiếu với ISO/IEC 25010:2023. Kiểm soát bảo mật tham khảo ISO/IEC 27002:2022 và ISO/IEC 27034-1:2011. C4 được dùng cho cấu trúc hệ thống; UML được dùng cho trạng thái và trình tự. Đây là cơ sở tổ chức tài liệu, không phải tuyên bố đã đạt chứng nhận.")
    base.add_table_caption(doc, 19, "Tiêu chuẩn và mô hình được áp dụng")
    base.add_table(doc, ["Nguồn", "Dùng để làm gì", "Giới hạn tuyên bố"], [
        ("ISO/IEC/IEEE 42010:2022", "Tổ chức architecture description theo stakeholder, concern, viewpoint và view.", "Chưa có đánh giá conformity theo điều khoản."),
        ("ISO/IEC/IEEE 29148:2018", "Giữ yêu cầu có mã, cần thiết, kiểm chứng được và truy vết.", "Báo cáo không thay thế SRS DOC-04."),
        ("ISO/IEC 25010:2023", "Nhắc đủ mối quan tâm về chất lượng, an toàn, vận hành và mở rộng.", "Chưa có số đo chất lượng production."),
        ("ISO/IEC 27002:2022 và ISO/IEC 27034-1:2011", "Tham khảo kiểm soát bảo mật và đưa bảo mật vào thiết kế ứng dụng.", "Không phải chứng nhận ISMS."),
        ("C4 và UML 2.5.1", "C4 mô tả cấu trúc; UML mô tả trạng thái và trình tự.", "Chỉ dùng loại sơ đồ trả lời câu hỏi cụ thể."),
    ], widths=[4.5, 8.3, 5.1], font_size=9.2)

    # Glossary
    doc.add_heading("Glossary Thuật ngữ dùng thống nhất", level=1)
    base.add_body_paragraph(doc, "Bảng dưới đây ghi cách hiểu thống nhất trong IDEA DDM. Thuật ngữ tiếng Anh được giữ lại khi đó là tên nghiệp vụ hoặc tên kỹ thuật quen thuộc.")
    glossary = [
        ("IDEA DDM", "Tên tạm thời của sản phẩm nội bộ quản lý tài liệu và dữ liệu thiết kế có kiểm soát."),
        ("Logical Document", "Danh tính ổn định của một tài liệu; không đổi chỉ vì đổi tên, chuyển folder, Check-in hoặc tạo Revision mới."),
        ("Artifact", "Nội dung file dạng byte được lưu trong kho file và nhận diện bằng digest."),
        ("Revision", "Mốc thay đổi nghiệp vụ lớn của cùng Logical Document, ví dụ A, B hoặc C."),
        ("Version", "Số bản nội dung trong một Revision; tăng khi Check-in có thay đổi và trở về 1 khi tạo Revision mới."),
        ("Generation", "Snapshot bất biến do hệ thống tạo để pin chính xác metadata, cấu trúc và Artifact; là mã kỹ thuật, không phải Version thứ hai cho người dùng."),
        ("Workspace", "Khu vực làm việc được quản lý trên máy người dùng, chứa file Checkout hoặc Reference và manifest phục hồi cục bộ."),
        ("Checkout", "Quyền giữ sửa một Logical Document cho đúng người, Workspace, expected Generation và thời hạn."),
        ("Reservation", "Bản ghi phía Server dùng để thực thi Checkout. Giao diện người dùng ưu tiên từ Checkout."),
        ("Reference", "Bản lấy để tham khảo; không có quyền Check-in vào Logical Document nguồn."),
        ("Check-in", "Thao tác có xác nhận để công bố thay đổi từ Workspace và kết thúc Checkout sau kết quả thành công hoặc No Change."),
        ("No Change", "Kết quả Check-in xác định nội dung có nghĩa không đổi; không tăng Version hoặc Generation nhưng vẫn kết thúc Checkout."),
        ("Working Head", "Generation hiện hành dùng làm cơ sở làm việc của một Revision chưa Release."),
        ("Expected Generation", "Generation mà Workspace dựa vào; dùng phát hiện bản làm việc đã cũ trước khi ghi."),
        ("Stale hoặc Out of date", "Bản làm việc dựa trên Generation cũ hơn Working Head; phải xử lý mà không ghi đè hoặc tự merge."),
        ("OperationId", "Mã ổn định của một thao tác; dùng tiếp tục hoặc tra lại đúng kết quả khi kết nối bị gián đoạn."),
        ("All or none", "Toàn bộ phạm vi được ghi thành công hoặc không tài liệu nào được công bố."),
        ("Review", "Quá trình xem xét một Generation và phạm vi đã pin theo workflow và approval policy."),
        ("Approve hoặc Reject", "Quyết định chấp thuận hoặc trả lại đúng bản đang review; Reject phải có lý do."),
        ("Release", "Quyết định phát hành chính thức một phạm vi sau khi kiểm tra lại đầy đủ điều kiện."),
        ("Release Record", "Bản ghi bất biến của một lần Release, pin đúng tài liệu, cấu trúc, policy, người, thời điểm và bằng chứng."),
        ("Controlled Release Package", "Gói có manifest, metadata, file, digest, cấu trúc và bằng chứng của một Release cụ thể."),
        ("Product Structure", "Dữ liệu cấu trúc cụm, chi tiết và quan hệ sản phẩm do IDEA DDM quản lý."),
        ("Structure Snapshot", "Ảnh chụp bất biến của Product Structure tại một mốc, pin đúng các Generation thành phần."),
        ("Representation", "Bản dẫn xuất như PDF hoặc preview, luôn gắn với đúng Generation nguồn và thông tin cách tạo."),
        ("RBAC", "Role-Based Access Control; kiểm soát quyền dựa trên Role Assignment."),
        ("Security Principal", "Người dùng hoặc Group có thể nhận Role Assignment."),
        ("Role Definition", "Định nghĩa có phiên bản chứa các Permission mà sản phẩm hỗ trợ."),
        ("Permission", "Mã hành động mà Server hiểu và kiểm tra, ví dụ Document.CheckIn."),
        ("Authorization Scope", "Phạm vi Role Assignment có hiệu lực: công ty, Project hoặc tài nguyên cụ thể."),
        ("Role Assignment", "Bản ghi gắn Principal, Role Definition phiên bản cụ thể và Scope."),
        ("Effective Permission", "Quyền hợp lệ tại thời điểm yêu cầu sau khi xét assignment trực tiếp, Group, Scope, trạng thái và thời hạn."),
        ("Business gate", "Điều kiện nghiệp vụ được kiểm tra sau RBAC, ví dụ trạng thái, chủ Checkout hoặc expected Generation."),
        ("Audit Evidence", "Bằng chứng chỉ bổ sung cho thao tác bảo vệ: người thực hiện, thời điểm, đối tượng, policy và kết quả."),
        ("Digest hoặc hash", "Dấu vân tay tính từ nội dung file, dùng kiểm tra toàn vẹn và nhận diện nội dung bất biến."),
        ("Session", "Phiên đăng nhập có thời hạn; có thể bị thu hồi khi tài khoản hoặc quyền thay đổi."),
        ("Fail closed", "Khi thiếu quyền, dữ liệu hoặc quy tắc thì từ chối an toàn thay vì tự bỏ qua điều kiện."),
        ("Backup", "Bản sao có kiểm soát của database, Artifact, cấu hình, policy và key cần để phục hồi."),
        ("Restore hoặc recovery", "Khôi phục một recovery set vào môi trường kiểm tra, đối soát dữ liệu rồi mới mở lại dịch vụ."),
        ("RPO", "Recovery Point Objective; lượng dữ liệu tối đa có thể mất tính theo thời gian. Giá trị cụ thể cần được phê duyệt và kiểm chứng."),
        ("RTO", "Recovery Time Objective; thời gian mục tiêu để khôi phục dịch vụ. Giá trị cụ thể cần được phê duyệt và kiểm chứng."),
        ("Qualification", "Hoạt động xác nhận môi trường, giới hạn và bằng chứng trước khi dùng thật."),
    ]
    base.add_table_caption(doc, 20, "Glossary thuật ngữ")
    base.add_table(doc, ["Thuật ngữ", "Cách hiểu trong IDEA DDM"], glossary, widths=[4.7, 13.2], font_size=9.0)

    # Sources
    doc.add_heading("Tài liệu nguồn và liên kết sơ đồ", level=1)
    base.add_body_paragraph(doc, "Các liên kết dưới đây mở tài liệu nguồn và sơ đồ chi tiết trong repository dự án. Khi gửi tài liệu ra ngoài cấu trúc thư mục hiện tại, cần gửi kèm thư mục evidence hoặc xuất sơ đồ thành một gói riêng.")
    for label, path in [
        ("DOC-03@0.7 Business Requirements", DOC03),
        ("DOC-04@0.13 Software Requirements Specification", DOC04),
        ("DOC-05@0.13 Architecture Description", DOC05),
        ("DOC-06@0.13 Data Integration and Migration Specification", DOC06),
        ("Thư viện 30 sơ đồ kiến trúc và dữ liệu", GALLERY),
        ("Standards applicability and version register", ROOT / "docs" / "governance" / "standards-register.md"),
    ]:
        p = doc.add_paragraph(style="List Bullet")
        base.add_hyperlink(p, label, path.resolve().as_uri())
    p = doc.add_paragraph()
    p.add_run("Trang chính thức dùng để xác nhận phiên bản tiêu chuẩn: ")
    for idx, (label, url) in enumerate([
        ("ISO/IEC/IEEE 42010:2022", "https://www.iso.org/standard/74393.html"),
        ("ISO/IEC/IEEE 29148:2018", "https://www.iso.org/standard/72089.html"),
        ("ISO/IEC 25010:2023", "https://www.iso.org/standard/78176.html"),
        ("ISO/IEC 27002:2022", "https://www.iso.org/standard/75652.html"),
    ]):
        if idx:
            p.add_run("; ")
        base.add_hyperlink(p, label, url)
    p.add_run(".")
    add_plain_note(doc, "Kết thúc tài liệu", "Mọi thay đổi làm ảnh hưởng Feature, Spec hoặc kiến trúc phải được ghi nhận tại tài liệu nguồn và review lại. Báo cáo này không phải nguồn quyết định độc lập.")

    _set_update_fields(doc)
    for style in doc.styles.element.findall(qn("w:style")):
        rpr = style.find(qn("w:rPr"))
        if rpr is not None:
            lang = rpr.find(qn("w:lang"))
            if lang is None:
                lang = OxmlElement("w:lang")
                rpr.append(lang)
            lang.set(qn("w:val"), "vi-VN")
            lang.set(qn("w:eastAsia"), "vi-VN")

    doc.save(str(OUTPUT))
    print(OUTPUT)


def main():
    figs = make_figures()
    build_document(figs)


if __name__ == "__main__":
    main()
