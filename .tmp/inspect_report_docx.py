from __future__ import annotations

import json
import re
import sys
import zipfile
from pathlib import Path
from html import unescape

from docx import Document
from lxml import etree


W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_R = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "pr": PKG_R}


def iter_block_text(doc: Document):
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            yield {"kind": "paragraph", "index": i, "style": paragraph.style.name, "text": text}
    for ti, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        yield {"kind": "table", "index": ti, "rows": rows}


def hyperlinks(docx_path: Path):
    with zipfile.ZipFile(docx_path) as zf:
        document = etree.fromstring(zf.read("word/document.xml"))
        rels = etree.fromstring(zf.read("word/_rels/document.xml.rels"))
        rel_targets = {
            rel.get("Id"): rel.get("Target")
            for rel in rels.findall(f"{{{PKG_R}}}Relationship")
            if rel.get("Type", "").endswith("/hyperlink")
        }
        out = []
        for h in document.xpath("//w:hyperlink", namespaces=NS):
            rid = h.get(f"{{{R}}}id")
            text = "".join(h.xpath(".//w:t/text()", namespaces=NS))
            out.append({"text": text, "target": rel_targets.get(rid), "anchor": h.get(f"{{{W}}}anchor")})
        return out


def main():
    source = Path(sys.argv[1])
    output = Path(sys.argv[2])
    doc = Document(source)
    payload = {
        "source": str(source),
        "paragraph_count": len(doc.paragraphs),
        "table_count": len(doc.tables),
        "sections": len(doc.sections),
        "blocks": list(iter_block_text(doc)),
        "hyperlinks": hyperlinks(source),
    }
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def svg_texts(directory: Path, names: list[str], output: Path):
    rows = {}
    for name in names:
        source = directory / name
        root = etree.fromstring(source.read_bytes())
        fragments = []
        for node in root.xpath("//*[local-name()='text' or local-name()='span' or local-name()='p']"):
            value = " ".join("".join(node.itertext()).split())
            value = unescape(value)
            if value and value not in fragments:
                fragments.append(value)
        rows[name] = fragments
    output.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "--svgs":
        svg_texts(Path(sys.argv[2]), sys.argv[3:-1], Path(sys.argv[-1]))
    else:
        main()
